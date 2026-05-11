import cv2
import pandas as pd
import os
import numpy as np
from pathlib import Path
from tqdm import tqdm

try:
    import openpyxl
    from openpyxl import load_workbook
except:
    openpyxl = None


def find_frame_data_file(folder_path):
    """pre 포함된 파일 (CSV 우선, 출력파일 제외)"""
    candidates = []
    exclude_stems = ['_trial_stats', '_final', '_pre_trial_stats', '_pre_final']

    # CSV 우선
    for file in folder_path.iterdir():
        if (file.name.startswith('._') or
                any(ex in file.stem.lower() for ex in exclude_stems)):
            continue
        if file.suffix.lower() == '.csv' and 'pre' in file.stem.lower():
            candidates.append(('csv', file))
            print(f"✅ FRAME DATA - pre CSV: {file.name}")

    # 엑셀
    if openpyxl and not candidates:
        for file in folder_path.iterdir():
            if (file.name.startswith('._') or
                    any(ex in file.stem.lower() for ex in exclude_stems)):
                continue
            if file.suffix.lower() in ['.xlsx', '.xls'] and 'pre' in file.stem.lower():
                candidates.append(('excel', file))
                print(f"✅ FRAME DATA - pre 엑셀: {file.name}")

    print(f"🔍 FRAME DATA 후보: {len(candidates)}개 → {candidates[0][1].name if candidates else '없음'}")
    return candidates[0] if candidates else None


def find_videos(folder_path):
    """비디오 찾기"""
    video_exts = ['.mp4', '.mkv', '.avi', '.mov', '.wmv', '.flv', '.webm', '.m4v']
    videos = [f for f in folder_path.iterdir()
              if f.suffix.lower() in video_exts and not f.name.startswith('._')]
    print(f"🎥 비디오: {len(videos)}개")
    return videos


def find_trial_file(folder_path):
    """동영상 파일과 정확히 같은 이름의 파일 (두 번째로 작은 동영상 기준)"""
    videos = find_videos(folder_path)
    if not videos:
        print("❌ 비디오 없음")
        return None

    # 용량 순으로 정렬 후 두 번째로 작은 동영상 선택
    videos_sorted = sorted(videos, key=lambda x: x.stat().st_size)
    if len(videos_sorted) < 2:
        print("⚠️ 비디오가 1개뿐 - 첫 번째 사용")
        video_path = videos_sorted[0]
    else:
        video_path = videos_sorted[1]  # 두 번째로 작은 것
        print(f"📹 두 번째로 작은 동영상 선택: {video_path.name} ({video_path.stat().st_size} bytes)")

    video_stem = video_path.stem

    print(f"🎥 선택 동영상: {video_path.name}")
    print(f"🔍 TRIAL 파일 찾기: '{video_stem}' 정확히 일치")

    # 동영상 stem과 정확히 일치하는 파일 (CSV 우선)
    for file in folder_path.iterdir():
        if (file.name.startswith('._') or file.stem != video_stem):
            continue

        if file.suffix.lower() == '.csv':
            print(f"✅ TRIAL - 정확히 일치 CSV: {file.name}")
            return ('csv', file)
        elif file.suffix.lower() in ['.xlsx', '.xls'] and openpyxl:
            print(f"✅ TRIAL - 정확히 일치 엑셀: {file.name}")
            return ('excel', file)

    print("❌ 동영상과 이름 일치하는 TRIAL 파일 없음")
    return None


def read_trial_file(file_info):
    """TRIAL 파일 읽기 (행수 체크 제거 - 모든 행 사용)"""
    file_type, file_path = file_info
    print(f"📖 TRIAL 로드: {file_path.name}")

    size_bytes = file_path.stat().st_size
    size_kb = size_bytes / 1024
    print(f"   📏 파일 크기: {size_bytes}bytes ({size_kb:.1f}KB)")

    # 엑셀 - 모든 행 읽기
    if file_type == 'excel' and openpyxl:
        try:
            wb = load_workbook(file_path, data_only=True)
            ws = wb.active
            data = [[cell.value for cell in row] for row in ws.iter_rows(min_row=2)]
            wb.close()
            df = pd.DataFrame(data)
            print(f"✅ 엑셀 로드: {len(df)}행")
        except Exception as e:
            print(f"❌ 엑셀 로드 실패: {e}")
            return None
    # CSV - 모든 행 읽기
    else:
        strategies = [
            {'encoding': 'utf-8-sig', 'sep': ',', 'header': None},
            {'encoding': 'utf-8', 'sep': ',', 'header': None},
            {'encoding': 'cp949', 'sep': ',', 'header': None},
            {'encoding': 'shift-jis', 'sep': ',', 'header': None},
        ]
        df = None
        for strategy in strategies:
            try:
                df = pd.read_csv(file_path, **strategy)
                if len(df) > 0:
                    print(f"✅ CSV 로드 성공 ({strategy['encoding']}) - {len(df)}행")
                    break
            except Exception as e:
                print(f"⚠️ {strategy.get('encoding', 'unknown')} 실패")
                continue

        if df is None:
            print("❌ 모든 CSV 인코딩 실패")
            return None

    if len(df) == 0:
        print("❌ TRIAL 데이터 없음")
        return None

    print(f"✅ TRIAL 확인: {len(df)}행 사용")
    print("📋 TRIAL 미리보기:")
    print(df.head().to_string())
    return df


def read_frame_data(file_info):
    """FRAME 데이터 읽기"""
    file_type, file_path = file_info
    if file_path.name.startswith('._'):
        return None

    if file_type == 'csv':
        return pd.read_csv(file_path)
    elif file_type == 'excel' and openpyxl:
        try:
            wb = load_workbook(file_path, data_only=True)
            ws = wb.active
            data = []
            for row in ws.iter_rows(values_only=True):
                data.append(row)
            wb.close()
            return pd.DataFrame(data[1:], columns=data[0])
        except:
            return None
    return None


def process_frame_data(frame_path_info, video_path, trial_df):
    """메인 처리 (Trial 시간 반올림 매칭)"""
    print("🚀 데이터 처리 시작!")

    df_frame = read_frame_data(frame_path_info)
    if df_frame is None or len(df_frame.columns) < 3:
        print("❌ FRAME 데이터 문제")
        return False

    print(f"📊 FRAME: {df_frame.shape}")
    col_frame_id, col_x, col_y = df_frame.columns[:3]

    # FPS 계산
    cap = cv2.VideoCapture(str(video_path))
    fps = cap.get(cv2.CAP_PROP_FPS)
    cap.release()
    print(f"📹 FPS: {fps}")

    # frame_id → 초 단위 시간
    df_frame['F'] = df_frame[col_frame_id].apply(
        lambda x: int(float(x)) / fps if pd.notna(x) and str(x).strip() != '' else None
    )
    df_frame['G'] = ""

    # 비교용 반올림 컬럼 (소수 1자리)
    df_frame['_time_round'] = df_frame['F'].round(1)

    stats_results = []

    for i in range(len(trial_df)):
        try:
            # trial 파일의 2,3열: 초 단위 시간
            start_time = float(trial_df.iloc[i, 1])
            end_time = float(trial_df.iloc[i, 2]) if len(trial_df.columns) > 2 else start_time + 10

            # trial 시간도 동일 자릿수로 반올림
            start_r = round(start_time, 1)
            end_r = round(end_time, 1)

            mask_start = df_frame['_time_round'] == start_r
            mask_end = df_frame['_time_round'] == end_r

            start_rows = df_frame[mask_start].index
            end_rows = df_frame[mask_end].index

            if len(start_rows) > 0 and len(end_rows) > 0:
                start_row, end_row = start_rows[0], end_rows[0]

                if end_row <= start_row:
                    print(f"⚠️ Trial{i + 1}: end_row <= start_row "
                          f"(start_row={start_row}, end_row={end_row})")
                    continue

                df_frame.loc[start_row, 'G'] = f"Trial{i + 1} start"
                df_frame.loc[end_row, 'G'] = f"Trial{i + 1} end"

                trial_data = df_frame.iloc[start_row:end_row + 1]
                x_data = trial_data[col_x].dropna()
                y_data = trial_data[col_y].dropna()

                if len(x_data) > 0:
                    stats_results.append({
                        'trial': i + 1,
                        'x_cor_aver': round(x_data.mean(), 1),
                        'y_cor_aver': round(y_data.mean(), 1),
                        'x_sd': round(x_data.std(), 1),
                        'y_sd': round(y_data.std(), 1),
                        'row_count': len(trial_data)
                    })
                    print(f"✅ Trial{i + 1}: {start_time}s~{end_time}s "
                          f"(rows {start_row}~{end_row}, n={len(trial_data)})")
            else:
                print(f"⚠️ Trial{i + 1}: 시간 매칭 실패 "
                      f"(start={start_time}, end={end_time}, "
                      f"start_rows={len(start_rows)}, end_rows={len(end_rows)})")
        except Exception as e:
            print(f"⚠️ Trial{i + 1} 에러: {e}")
            continue

    # 기존 파일 무조건 덮어쓰기 (pre 파일명 사용)
    frame_path = frame_path_info[1]
    stats_file = frame_path.parent / (frame_path.stem + "_pre_trial_stats.csv")
    final_file = frame_path.parent / (frame_path.stem + "_pre_final.csv")

    if stats_file.exists():
        print(f"🗑️ 기존 통계 삭제: {stats_file}")
        stats_file.unlink()
    if final_file.exists():
        print(f"🗑️ 기존 최종 삭제: {final_file}")
        final_file.unlink()

    if stats_results:
        pd.DataFrame(stats_results).to_csv(stats_file, index=False, encoding='utf-8-sig')
        print(f"📈 통계 저장: {stats_file}")

    df_frame.to_csv(final_file, index=False, encoding='utf-8-sig')
    print(f"✅ 최종 저장: {final_file}")
    return True


def process_folder(folder_path):
    """폴더 단위 처리"""
    folder = Path(folder_path)

    frame_info = find_frame_data_file(folder)
    if not frame_info:
        print(f"❌ {folder.name}: FRAME DATA 없음")
        return False

    videos = find_videos(folder)
    if len(videos) == 0:
        print(f"❌ {folder.name}: 비디오 없음")
        return False

    # 두 번째로 작은 동영상 선택
    videos_sorted = sorted(videos, key=lambda x: x.stat().st_size)
    if len(videos_sorted) < 2:
        print(f"⚠️ {folder.name}: 비디오가 1개뿐 - 첫 번째 사용")
        video_path = videos_sorted[0]
    else:
        video_path = videos_sorted[1]

    trial_info = find_trial_file(folder)
    if not trial_info:
        print(f"❌ {folder.name}: TRIAL 파일 없음")
        return False

    trial_df = read_trial_file(trial_info)
    if trial_df is None:
        print(f"❌ {folder.name}: TRIAL 로드 실패")
        return False

    print(f"🎯 {folder.name} 처리! (TRIAL: {len(trial_df)}행)")
    return process_frame_data(frame_info, video_path, trial_df)


def batch_process(root_folder):
    """배치 처리"""
    root = Path(root_folder)
    folders = [f for f in root.rglob('*') if f.is_dir()]

    print(f"📁 총 폴더: {len(folders)}개")
    success = 0

    for folder in tqdm(folders, desc="처리중"):
        try:
            if process_folder(folder):
                success += 1
        except Exception as e:
            print(f"❌ {folder.name}: {e}")

    print(f"\n🎉 완료! 성공: {success}/{len(folders)}")


if __name__ == "__main__":
    ROOT_FOLDER = r'/Volumes/ボリューム/2025_gaze_experiment/sub13'
    batch_process(ROOT_FOLDER)